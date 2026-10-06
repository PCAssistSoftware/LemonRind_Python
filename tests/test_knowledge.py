"""Tests for knowledge bases: chunking, reading files and web pages, storage, ingestion, and the module."""

from __future__ import annotations

from pathlib import Path

import httpx2 as httpx
import pytest

from lemonrind.chats import ChatRepository, Conversation
from lemonrind.config import Settings
from lemonrind.embeddings import Embedder
from lemonrind.modules.base import ChatContext
from lemonrind.modules.knowledge import (
    OPTION_KEY,
    DuplicateNameError,
    KnowledgeIngestor,
    KnowledgeModule,
    KnowledgeRepository,
)
from lemonrind.modules.knowledge.chunker import chunk_text
from lemonrind.modules.knowledge.extract import ExtractionError, extract_text, is_supported
from lemonrind.modules.knowledge.webpage import WebPageError, fetch_page_text, html_to_text
from lemonrind.storage import Database
from tests.test_memory import FakeMemoryClient

# --- chunking -----------------------------------------------------------------------------------------------


def numbered(count: int) -> str:
    return " ".join(f"w{n}" for n in range(count))


def test_short_text_is_one_chunk_and_empty_text_is_none():
    assert chunk_text("just a few words", 400, 60) == ["just a few words"]
    assert chunk_text("   \n  ", 400, 60) == []
    assert chunk_text("", 400, 60) == []


def test_chunks_overlap_by_the_requested_number_of_words():
    chunks = chunk_text(numbered(25), words_per_chunk=10, overlap=3)

    assert chunks[0].split()[0] == "w0" and chunks[0].split()[-1] == "w9"
    assert chunks[1].split()[0] == "w7"  # starts 3 words before chunk 0 ended
    assert chunks[1].split()[-1] == "w16"
    assert chunks[-1].split()[-1] == "w24"  # nothing is lost at the end
    assert len(chunks) == 4


def test_no_extra_chunk_that_only_repeats_the_tail():
    # exactly one chunk's worth, and exactly two (with overlap): no trailing duplicate
    assert len(chunk_text(numbered(10), 10, 3)) == 1
    assert len(chunk_text(numbered(17), 10, 3)) == 2
    assert chunk_text(numbered(17), 10, 3)[-1].split()[-1] == "w16"


def test_line_breaks_inside_a_chunk_are_preserved():
    text = "first line\n\nsecond paragraph here"
    assert chunk_text(text, 400, 60) == [text]


@pytest.mark.parametrize(("size", "overlap"), [(0, 0), (10, 10), (10, -1)])
def test_bad_chunk_settings_are_rejected(size, overlap):
    with pytest.raises(ValueError):
        chunk_text("some text", size, overlap)


# --- reading files ------------------------------------------------------------------------------------------


def make_pdf(text: str) -> bytes:
    """A minimal real PDF (with a proper cross-reference table) holding one line of text."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<</Root 1 0 R/Size %d>>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return out


def test_text_files_are_read_whatever_their_encoding(tmp_path: Path):
    (tmp_path / "a.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "b.md").write_bytes(
        b"\xef\xbb\xbfwith marker"
    )  # utf-8 with the Windows byte-order mark
    (tmp_path / "c.csv").write_bytes("caf\xe9".encode("latin-1"))  # not valid utf-8

    assert extract_text(tmp_path / "a.txt") == "hello"
    assert extract_text(tmp_path / "b.md") == "with marker"
    assert extract_text(tmp_path / "c.csv") == "café"


def test_pdf_text_is_extracted_and_a_scanned_pdf_is_explained(tmp_path: Path):
    from pypdf import PdfWriter

    (tmp_path / "doc.pdf").write_bytes(make_pdf("Hello pizza world"))
    writer = PdfWriter()
    writer.add_blank_page(100, 100)  # a page with no text layer, like a scan
    with (tmp_path / "scan.pdf").open("wb") as handle:
        writer.write(handle)

    assert extract_text(tmp_path / "doc.pdf") == "Hello pizza world"
    with pytest.raises(ExtractionError, match="no text layer"):
        extract_text(tmp_path / "scan.pdf")


def test_word_documents_give_paragraphs_and_tables(tmp_path: Path):
    from docx import Document

    document = Document()
    document.add_paragraph("Holiday policy")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Days"
    table.rows[0].cells[1].text = "25"
    document.save(str(tmp_path / "a.docx"))

    text = extract_text(tmp_path / "a.docx")

    assert "Holiday policy" in text and "Days\t25" in text


def test_excel_sheets_become_tab_separated_lines(tmp_path: Path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Item", "Price"])
    sheet.append(["Mileage", 0.45])
    workbook.save(tmp_path / "a.xlsx")

    assert extract_text(tmp_path / "a.xlsx") == "Sheet: Rates\nItem\tPrice\nMileage\t0.45"


def test_unreadable_files_raise_a_readable_error(tmp_path: Path):
    (tmp_path / "bad.pdf").write_bytes(b"this is not a pdf")
    (tmp_path / "empty.txt").write_text("   \n", encoding="utf-8")
    (tmp_path / "a.exe").write_bytes(b"MZ")

    with pytest.raises(ExtractionError, match="Could not read 'bad.pdf'"):
        extract_text(tmp_path / "bad.pdf")
    with pytest.raises(ExtractionError, match="no readable text"):
        extract_text(tmp_path / "empty.txt")
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_text(tmp_path / "a.exe")
    assert is_supported(Path("x.PDF")) and not is_supported(Path("x.exe"))


# --- web pages ----------------------------------------------------------------------------------------------

PAGE = """<html><head><title> Guide </title><style>p {color: red}</style></head>
<body><script>alert('x')</script><h1>Refunds</h1><p>Refunds take   five days.</p>
<ul><li>First &amp; second</li></ul><noscript>enable js</noscript></body></html>"""


def test_html_is_reduced_to_its_readable_text():
    title, text = html_to_text(PAGE)

    assert title == "Guide"
    assert text == "Refunds\nRefunds take five days.\nFirst & second"
    assert "alert" not in text and "color" not in text and "enable js" not in text


async def test_a_page_is_fetched_and_read():
    transport = httpx.MockTransport(
        lambda request: httpx.Response(200, text=PAGE, headers={"content-type": "text/html"})
    )

    title, text = await fetch_page_text("https://example.com/guide", transport=transport)

    assert title == "Guide" and "Refunds take five days." in text


@pytest.mark.parametrize(
    ("url", "response", "message"),
    [
        ("ftp://example.com/x", httpx.Response(200, text="x"), "http://"),
        ("https://example.com", httpx.Response(404), "Could not fetch"),
        (
            "https://example.com",
            httpx.Response(200, content=b"%PDF", headers={"content-type": "application/pdf"}),
            "Cannot read content",
        ),
        (
            "https://example.com",
            httpx.Response(
                200, text="<html><script>x()</script></html>", headers={"content-type": "text/html"}
            ),
            "no readable text",
        ),
    ],
)
async def test_bad_pages_give_clear_errors(url, response, message):
    transport = httpx.MockTransport(lambda request: response)
    with pytest.raises(WebPageError, match=message):
        await fetch_page_text(url, transport=transport)


# --- the store -------------------------------------------------------------------------------------------------


KB_FOLDER = Path()  # set for every test in this file by the fixture below


@pytest.fixture(autouse=True)
def _knowledge_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Each test gets its own empty folder for knowledge base files."""
    monkeypatch.setitem(globals(), "KB_FOLDER", tmp_path / "knowledge")


def make_repo() -> KnowledgeRepository:
    return KnowledgeRepository(KB_FOLDER)


def test_knowledge_base_names_are_unique_ignoring_case():
    repo = make_repo()
    repo.create_kb("Work")

    with pytest.raises(DuplicateNameError):
        repo.create_kb("work")
    with pytest.raises(ValueError):
        repo.create_kb("   ")
    assert repo.find_kb("WORK") is not None and [kb.name for kb in repo.list_kbs()] == ["Work"]


def test_only_chunks_of_ready_sources_are_searchable_and_failure_discards_chunks():
    repo = make_repo()
    kb = repo.create_kb("Work")
    ready = repo.create_source(kb.id, "text", None, "ready.txt")
    pending = repo.create_source(kb.id, "text", None, "pending.txt")
    repo.add_chunks(kb.id, ready.id, 0, ["a"], [[1, 0]])
    repo.add_chunks(kb.id, pending.id, 0, ["b"], [[0, 1]])
    repo.mark_ready(ready.id, 1)

    assert [c.content for c, _ in repo.searchable_chunks(kb.id)] == ["a"]  # "b" is still ingesting

    repo.mark_failed(ready.id, "boom")
    failed = repo.get_source(ready.id)
    assert failed is not None and (failed.status, failed.error, failed.chunk_count) == (
        "failed",
        "boom",
        0,
    )
    assert repo.searchable_chunks(kb.id) == []


def test_sources_cut_off_by_a_restart_are_marked_failed():
    repo = make_repo()
    kb = repo.create_kb("Work")
    source = repo.create_source(kb.id, "file", "a.txt", "a.txt")
    repo.add_chunks(kb.id, source.id, 0, ["x"], [[1, 0]])

    assert repo.fail_interrupted() == 1

    after = repo.get_source(source.id)
    assert after is not None and after.status == "failed" and "Interrupted" in (after.error or "")


def test_deleting_a_knowledge_base_or_source_removes_what_hangs_off_it():
    repo = make_repo()
    kb = repo.create_kb("Work")
    keep = repo.create_source(kb.id, "text", None, "keep")
    drop = repo.create_source(kb.id, "text", None, "drop")
    for source in (keep, drop):
        repo.add_chunks(kb.id, source.id, 0, [source.display_name], [[1, 0]])
        repo.mark_ready(source.id, 1)

    repo.delete_source(drop.id)
    assert [c.content for c, _ in repo.searchable_chunks(kb.id)] == ["keep"]

    repo.delete_kb(kb.id)
    assert repo.list_sources(kb.id) == [] and repo.searchable_chunks(kb.id) == []


def test_chunk_labels_name_the_file_inside_a_folder_source():
    repo = make_repo()
    kb = repo.create_kb("Work")
    source = repo.create_source(kb.id, "folder", "docs", "docs")
    repo.add_chunks(kb.id, source.id, 0, ["x"], [[1, 0]], label="sub/a.md")
    repo.mark_ready(source.id, 1)

    assert repo.file_labels(source.id) == ["sub/a.md"]
    assert repo.searchable_chunks(kb.id)[0][0].source_name == "sub/a.md"


def test_per_chat_options_are_stored_and_removed():
    chats = ChatRepository(Database(":memory:"))
    session = chats.create_session()

    chats.set_option(session.id, "knowledge_base", "kb-1")
    chats.set_option(session.id, "other", "x")
    chats.set_option(session.id, "other", None)

    stored = chats.get_session(session.id)
    assert stored is not None and dict(stored.options) == {"knowledge_base": "kb-1"}


# --- ingestion -----------------------------------------------------------------------------------------------


def make_world(client: FakeMemoryClient | None = None, *, transport=None):
    settings = Settings()
    settings.modules.knowledge.chunk_words = 50
    settings.modules.knowledge.chunk_overlap = 10
    settings.modules.knowledge.min_similarity = (
        0.6  # the fake embedding scores related texts 1.0, unrelated 0.5
    )
    repo = make_repo()
    embedder = Embedder(settings, client or FakeMemoryClient())  # type: ignore[arg-type]
    ingestor = KnowledgeIngestor(settings, repo, embedder, transport=transport)
    module = KnowledgeModule(settings, repo, embedder, ingestor)
    return settings, repo, ingestor, module


async def test_a_file_is_chunked_embedded_and_made_searchable(tmp_path: Path):
    _, repo, ingestor, _ = make_world()
    kb = repo.create_kb("Work")
    file = tmp_path / "menu.txt"
    file.write_text(
        "pizza food " * 60, encoding="utf-8"
    )  # 120 words: three chunks of 50 with overlap 10

    source = await ingestor.add_file(kb.id, file)

    assert (source.status, source.chunk_count, source.display_name) == ("ready", 3, "menu.txt")
    assert len(repo.searchable_chunks(kb.id)) == 3


async def test_an_unreadable_file_is_recorded_as_failed_not_raised(tmp_path: Path):
    _, repo, ingestor, _ = make_world()
    kb = repo.create_kb("Work")
    (tmp_path / "bad.pdf").write_bytes(b"nope")

    source = await ingestor.add_file(kb.id, tmp_path / "bad.pdf")

    assert source.status == "failed" and "Could not read 'bad.pdf'" in (source.error or "")


async def test_without_an_embedding_model_ingestion_fails_with_the_reason(tmp_path: Path):
    _, repo, ingestor, _ = make_world(FakeMemoryClient(with_embedding_model=False))
    kb = repo.create_kb("Work")

    source = await ingestor.add_text(kb.id, "note", "pizza")

    assert source.status == "failed" and "no embedding model" in (source.error or "").lower()
    assert repo.searchable_chunks(kb.id) == []


async def test_a_folder_is_one_source_and_one_bad_file_does_not_sink_the_rest(tmp_path: Path):
    _, repo, ingestor, _ = make_world()
    kb = repo.create_kb("Work")
    folder = tmp_path / "docs"
    (folder / "sub").mkdir(parents=True)
    (folder / "a.txt").write_text("pizza food", encoding="utf-8")
    (folder / "sub" / "b.md").write_text("manchester city", encoding="utf-8")
    (folder / "bad.pdf").write_bytes(b"nope")  # supported type but unreadable: skipped
    (folder / "ignored.exe").write_bytes(b"MZ")  # unsupported type: not even tried

    source = await ingestor.add_folder(kb.id, folder)

    assert (source.status, source.chunk_count, source.source_type) == ("ready", 2, "folder")
    assert repo.file_labels(source.id) == ["a.txt", "sub/b.md"]
    assert len(repo.list_sources(kb.id)) == 1


async def test_empty_or_unreadable_folders_fail_with_a_message(tmp_path: Path):
    _, repo, ingestor, _ = make_world()
    kb = repo.create_kb("Work")
    empty = tmp_path / "empty"
    empty.mkdir()
    broken = tmp_path / "broken"
    broken.mkdir()
    (broken / "bad.pdf").write_bytes(b"nope")

    no_files = await ingestor.add_folder(kb.id, empty)
    none_readable = await ingestor.add_folder(kb.id, broken)

    assert "No supported files" in (no_files.error or "")
    assert "None of the files" in (none_readable.error or "")


async def test_a_website_and_pasted_text_are_ingested():
    transport = httpx.MockTransport(
        lambda request: httpx.Response(
            200,
            text="<html><title>Menu</title><p>pizza food</p></html>",
            headers={"content-type": "text/html"},
        )
    )
    _, repo, ingestor, _ = make_world(transport=transport)
    kb = repo.create_kb("Work")

    page = await ingestor.add_website(kb.id, "https://example.com/menu")
    note = await ingestor.add_text(kb.id, "", "manchester city")
    dead = await ingestor.add_website(kb.id, "ftp://nope")

    assert (page.status, page.source_type) == ("ready", "website")
    assert (note.status, note.display_name) == ("ready", "Pasted text")
    assert dead.status == "failed" and "http://" in (dead.error or "")


# --- the module -----------------------------------------------------------------------------------------------


def chat_with(kb_id: str | None) -> ChatContext:
    return ChatContext(options={OPTION_KEY: kb_id} if kb_id else {})


async def attached_world(tmp_path: Path):
    _, repo, ingestor, module = make_world()
    kb = repo.create_kb("Handbook")
    await ingestor.add_text(kb.id, "menu.txt", "The canteen serves pizza food on Fridays.")
    await ingestor.add_text(kb.id, "map.txt", "The office is in Manchester city centre.")
    return module, repo, kb


async def test_relevant_chunks_and_the_source_list_are_added_for_an_attached_knowledge_base(
    tmp_path: Path,
):
    module, _, kb = await attached_world(tmp_path)

    text = await module.turn_context("What pizza food is there?", chat=chat_with(kb.id))

    assert "knowledge base called 'Handbook'" in text
    assert "- menu.txt" in text and "- map.txt" in text  # the authoritative source list
    assert "[Source: menu.txt]\nThe canteen serves pizza food on Fridays." in text
    assert (
        "Manchester" not in text.split("Relevant excerpts")[1]
    )  # the unrelated chunk is not included


async def test_an_unrelated_question_still_gets_the_source_list_but_no_excerpts(tmp_path: Path):
    module, _, kb = await attached_world(tmp_path)

    text = await module.turn_context("Tell me a joke.", chat=chat_with(kb.id))

    assert "- menu.txt" in text and "Relevant excerpts" not in text


async def test_folder_sources_list_the_files_inside_them(tmp_path: Path):
    _, repo, ingestor, module = make_world()
    kb = repo.create_kb("Docs")
    folder = tmp_path / "policies"
    folder.mkdir()
    (folder / "holiday.txt").write_text("pizza food", encoding="utf-8")
    await ingestor.add_folder(kb.id, folder)

    text = await module.turn_context("hello", chat=chat_with(kb.id))

    assert "- policies (folder containing: holiday.txt)" in text


async def test_without_an_attached_knowledge_base_every_message_says_so_whatever_it_is_about():
    _, _, _, module = make_world()

    for message in ("hello", "What is in the knowledge base?", "Summarise my notes"):
        note = await module.turn_context(message, chat=chat_with(None))
        assert "No knowledge base is attached" in note and "file system" in note
    vanished = await module.turn_context("hello", chat=chat_with("deleted-id"))
    assert "No knowledge base is attached" in vanished  # a vanished one counts as none


async def test_an_attached_knowledge_base_with_no_ready_sources_says_so():
    _, repo, _, module = make_world()
    kb = repo.create_kb("Empty")

    text = await module.turn_context("hello", chat=chat_with(kb.id))

    assert "no sources ready yet" in text


async def test_startup_fails_sources_left_ingesting():
    _, repo, _, module = make_world()
    kb = repo.create_kb("Work")
    source = repo.create_source(kb.id, "file", "a", "a")

    await module.on_startup()

    after = repo.get_source(source.id)
    assert after is not None and after.status == "failed"


# --- per-chat options through a conversation -----------------------------------------------------------------


class Quiet:
    async def stream_chat(self, messages, model, **kwargs):
        from lemonrind.lemonade import Finished, TextDelta
        from lemonrind.lemonade.events import RequestStats

        yield TextDelta("ok")
        yield Finished(RequestStats(), "stop")


class Recorder:
    def __init__(self) -> None:
        self.chats: list[ChatContext] = []

    async def stable_context(self) -> str:
        return ""

    async def turn_context(self, user_text: str, *, chat: ChatContext) -> str:
        self.chats.append(chat)
        return ""

    async def after_turn(self, user_text: str, reply_text: str, *, model: str) -> None:
        pass


async def test_chat_options_reach_modules_are_saved_with_the_chat_and_restored_on_open():
    chats = ChatRepository(Database(":memory:"))
    recorder = Recorder()
    conversation = Conversation(
        client=Quiet(), repo=chats, system_prompt="x", model="m", context=recorder
    )

    conversation.set_option(OPTION_KEY, "kb-1")  # chosen before the chat exists
    await conversation.send("hello")

    assert recorder.chats[-1].options == {OPTION_KEY: "kb-1"}
    session = conversation.session
    assert session is not None and dict(session.options) == {
        OPTION_KEY: "kb-1"
    }  # saved at first reply

    conversation.set_option(OPTION_KEY, "kb-2")  # changed afterwards: saved immediately
    reopened = Conversation(client=Quiet(), repo=chats, system_prompt="x", model="m")
    reopened.open(chats.get_session(session.id))  # type: ignore[arg-type]
    assert reopened.options == {OPTION_KEY: "kb-2"}

    conversation.new()
    assert conversation.options == {}  # a new chat starts with nothing attached
