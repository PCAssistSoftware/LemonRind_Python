"""Getting plain text out of files: PDF, Word, Excel, and plain text.

Each format needs its own library, because a ``.docx`` or ``.pdf`` is not text but a structured file:

* **PDF** - ``pypdf`` reads the text layer of each page. A *scanned* PDF is just pictures of pages and has
  no text layer; there is nothing to read without OCR (optical character recognition), so we say so
  instead of silently returning nothing (which would make the model, and you, wonder what happened).
* **Word (.docx)** - ``python-docx`` gives paragraphs and tables.
* **Excel (.xlsx)** - ``openpyxl`` gives cells; each sheet becomes tab-separated lines.
* **Text** (.txt, .md, .csv) - read directly.

The libraries are imported inside the functions that use them. They are slow to import and only needed when
that kind of file is actually read, so start-up stays quick.

Python ideas used here:

* ``pathlib.Path`` methods (``suffix``, ``read_bytes``).
* A ``dict`` of functions as a dispatch table instead of a long ``if/elif`` chain.
* Wrapping third-party exceptions in our own (``ExtractionError``) so callers handle one kind of failure.
"""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path


class ExtractionError(Exception):
    """A file could not be read as text. The message is fit to show the user."""


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    pages = [(page.extract_text() or "").strip() for page in PdfReader(path).pages]
    text = "\n\n".join(page for page in pages if page)
    if not text:
        raise ExtractionError(
            f"'{path.name}' has no text layer (it is probably a scanned image). Text cannot be read "
            "from it without OCR."
        )
    return text


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(
        path, read_only=True, data_only=True
    )  # data_only: values, not formulas
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows = [
                "\t".join("" if cell is None else str(cell) for cell in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sheets.append(f"Sheet: {sheet.title}\n" + "\n".join(row for row in rows if row.strip()))
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode(
            "utf-8-sig"
        )  # "utf-8-sig" also removes the invisible marker Windows adds
    except UnicodeDecodeError:
        return data.decode("latin-1")  # never fails: every byte is a valid latin-1 character


_READERS: dict[str, Callable[[Path], str]] = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
    ".txt": _read_text,
    ".md": _read_text,
    ".csv": _read_text,
}

SUPPORTED_EXTENSIONS = tuple(_READERS)


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in _READERS


def extract_text(path: Path) -> str:
    """The text of a supported file. Raises ``ExtractionError`` with a readable message otherwise."""
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise ExtractionError(
            f"Unsupported file type '{path.suffix}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    try:
        text = reader(path)
    except ExtractionError:
        raise
    except (
        Exception
    ) as error:  # a corrupt or password-protected file raises library-specific errors
        raise ExtractionError(
            f"Could not read '{path.name}': {type(error).__name__}: {error}"
        ) from error
    if not text.strip():
        raise ExtractionError(f"'{path.name}' contains no readable text.")
    return text
